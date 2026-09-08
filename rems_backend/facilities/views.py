

from datetime import timedelta
from decimal import Decimal, InvalidOperation

from django.contrib.auth import get_user_model
from django.db import transaction
from django.utils import timezone

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from rest_framework import status
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.viewsets import ModelViewSet

from accounts.permissions import (
    IsAdmin,
    IsAdminOrResident,
)

from notifications.models import Notification
from notifications.services import NotificationService

from .models import (
    Facility,
    Booking,
    BookingGuest,
    BookingPayment,
    BookingInspection,
)

from .serializers import (
    FacilitySerializer,
    BookingSerializer,
    BookingAvailabilitySerializer,
    BookingGuestSerializer,
    BookingPaymentSerializer,
    BookingInspectionSerializer,
)

from .services import FacilityBookingService


User = get_user_model()


# ============================================================
# FACILITIES
# ============================================================

class FacilityViewSet(ModelViewSet):
    """
    Facility catalog and administration.

    GET
        /api/facilities/
        /api/facilities/active/
        /api/facilities/available/

    ADMIN
        Full CRUD.
    """

    queryset = Facility.objects.all()

    serializer_class = FacilitySerializer

    permission_classes = [
        IsAuthenticated,
        IsAdminOrResident,
    ]

    search_fields = [
        "name",
        "description",
        "location",
    ]

    filterset_fields = [
        "facility_type",
        "is_active",
        "is_bookable",
        "requires_approval",
    ]

    ordering_fields = [
        "name",
        "facility_type",
        "rental_fee",
        "created_at",
        "updated_at",
    ]

    ordering = [
        "name",
    ]

    def get_permissions(self):
        if self.action in [
            "create",
            "update",
            "partial_update",
            "destroy",
        ]:
            permission_classes = [
                IsAuthenticated,
                IsAdmin,
            ]
        else:
            permission_classes = [
                IsAuthenticated,
                IsAdminOrResident,
            ]

        return [
            permission()
            for permission in permission_classes
        ]

    # ========================================================
    # ACTIVE FACILITIES
    # ========================================================

    @action(
        detail=False,
        methods=["get"],
        url_path="active",
    )
    def active(self, request):
        facilities = (
            self.get_queryset()
            .filter(
                is_active=True,
                is_bookable=True,
            )
        )

        serializer = self.get_serializer(
            facilities,
            many=True,
        )

        return Response(serializer.data)

    # ========================================================
    # AVAILABILITY
    # ========================================================

    @action(
        detail=False,
        methods=["get"],
        url_path="available",
    )
    def available(self, request):
        serializer = BookingAvailabilitySerializer(
            data=request.query_params
        )

        serializer.is_valid(
            raise_exception=True
        )

        data = serializer.validated_data

        facility = data["facility"]

        availability = (
            FacilityBookingService.check_availability(
                facility=facility,
                booking_date=data["booking_date"],
                start_time=data["start_time"],
                end_time=data["end_time"],
                estimated_guests=data.get(
                    "estimated_guests",
                    1,
                ),
            )
        )

        financials = (
            FacilityBookingService.calculate_financials(
                facility
            )
        )

        return Response(
            {
                "facility": facility.id,
                "facility_name": facility.name,
                "booking_date": data["booking_date"],
                "start_time": data["start_time"],
                "end_time": data["end_time"],
                **availability,
                **financials,
            }
        )


# ============================================================
# BOOKINGS
# ============================================================

class BookingViewSet(ModelViewSet):
    """
    Main facility booking workflow.

    Resident:
        PENCIL
            -> PENDING
            -> APPROVED
            -> IN_USE
            -> INSPECTION_PENDING
            -> REFUND_PENDING
            -> CLOSED

    Alternative terminal states:
        REJECTED
        CANCELLED

    IMPORTANT:
        GET requests never mutate booking state.
    """

    queryset = (
        Booking.objects
        .select_related(
            "facility",
            "resident",
            "resident__user",
            "approved_by",
            "security_cleared_by",
            "inspection",
        )
        .prefetch_related(
            "guests",
            "payments",
        )
        .all()
    )

    serializer_class = BookingSerializer

    permission_classes = [
        IsAuthenticated,
        IsAdminOrResident,
    ]

    http_method_names = [
        "get",
        "post",
        "patch",
        "delete",
        "head",
        "options",
    ]

    search_fields = [
        "facility__name",
        "resident__first_name",
        "resident__middle_name",
        "resident__last_name",
        "event_description",
        "supplier_details",
    ]

    filterset_fields = [
        "facility",
        "resident",
        "status",
        "booking_date",
        "event_type",
        "security_clearance_completed",
        "inspection_required",
    ]

    ordering_fields = [
        "booking_date",
        "start_time",
        "end_time",
        "created_at",
        "updated_at",
        "submitted_at",
    ]

    ordering = [
        "-booking_date",
        "-start_time",
    ]

    # ========================================================
    # QUERYSET VISIBILITY
    # ========================================================

    def get_queryset(self):
        queryset = super().get_queryset()

        user = self.request.user

        if user.role == user.Roles.ADMIN:
            return queryset

        return queryset.filter(
            resident__user=user
        )

    # ========================================================
    # RESIDENT LOOKUP
    # ========================================================

    def _get_resident(self):
        user = self.request.user

        resident = getattr(
            user,
            "resident",
            None,
        )

        if resident:
            return resident

        from residents.models import Resident

        try:
            return (
                Resident.objects
                .select_related("user")
                .get(user=user)
            )
        except Resident.DoesNotExist:
            return None

    # ========================================================
    # NOTIFICATION URL
    # ========================================================

    @staticmethod
    def _resident_booking_url(resident):
        role = str(
            getattr(
                resident.user,
                "role",
                "",
            )
        ).strip().upper()

        if role == "TENANT":
            return "/tenant/facilities/bookings"

        return "/homeowner/facilities/bookings"

    # ========================================================
    # NOTIFICATION DISPATCH
    # ========================================================

    @staticmethod
    def _send_notification(
        *,
        recipient,
        actor=None,
        title,
        message,
        action_url="",
        metadata=None,
        priority=Notification.Priority.INFO,
    ):
        """
        Resolve the notification service method without making
        facility bookings depend on one hard-coded notification
        category method.
        """

        notifier = (
            getattr(
                NotificationService,
                "facility",
                None,
            )
            or getattr(
                NotificationService,
                "property",
                None,
            )
            or getattr(
                NotificationService,
                "create",
                None,
            )
        )

        if not notifier:
            return

        notifier(
            recipient=recipient,
            actor=actor,
            title=title,
            message=message,
            action_url=action_url,
            metadata=metadata or {},
            priority=priority,
        )

    # ========================================================
    # NOTIFY ADMINS
    # ========================================================

    @staticmethod
    def _notify_admins(
        *,
        title,
        message,
        actor=None,
        action_url="/admin/facilities/bookings",
        metadata=None,
        priority=Notification.Priority.INFO,
    ):
        admins = (
            User.objects
            .filter(
                role=User.Roles.ADMIN,
                is_active=True,
            )
        )

        for admin_user in admins:
            BookingViewSet._send_notification(
                recipient=admin_user,
                actor=actor,
                title=title,
                message=message,
                action_url=action_url,
                metadata=metadata,
                priority=priority,
            )

    # ========================================================
    # NOTIFY RESIDENT
    # ========================================================

    @staticmethod
    def _notify_resident(
        *,
        booking,
        title,
        message,
        actor=None,
        metadata=None,
        priority=Notification.Priority.INFO,
    ):
        if not booking.resident:
            return

        resident_user = booking.resident.user

        BookingViewSet._send_notification(
            recipient=resident_user,
            actor=actor,
            title=title,
            message=message,
            action_url=(
                BookingViewSet._resident_booking_url(
                    booking.resident
                )
            ),
            metadata=metadata,
            priority=priority,
        )

    # ========================================================
    # CREATE
    # ========================================================

    def perform_create(self, serializer):
        user = self.request.user

        # ----------------------------------------------------
        # ADMIN
        # ----------------------------------------------------

        if user.role == user.Roles.ADMIN:
            booking = serializer.save(
                status=Booking.Status.PENDING,
                pencil_created_at=None,
                pencil_expires_at=None,
            )

            financials = (
                FacilityBookingService.apply_financials(
                    booking
                )
            )

            booking.rental_fee = financials[
                "rental_fee"
            ]

            booking.discount_amount = financials[
                "discount_amount"
            ]

            booking.security_deposit = financials[
                "security_deposit"
            ]

            booking.total_amount_due = financials[
                "total_amount_due"
            ]

            booking.save(
                update_fields=[
                    "rental_fee",
                    "discount_amount",
                    "security_deposit",
                    "total_amount_due",
                    "updated_at",
                ]
            )

            return

        # ----------------------------------------------------
        # RESIDENT
        # ----------------------------------------------------

        resident = self._get_resident()

        if not resident:
            raise ValidationError(
                {
                    "resident":
                        "No resident profile is associated with your account."
                }
            )

        now = timezone.now()

        booking = serializer.save(
            resident=resident,
            status=Booking.Status.PENCIL,
            pencil_created_at=now,
            pencil_expires_at=(
                now
                + timedelta(
                    minutes=(
                        FacilityBookingService
                        .PENCIL_HOLD_MINUTES
                    )
                )
            ),
        )

        # ----------------------------------------------------
        # IMPORTANT
        #
        # Store the financial snapshot immediately so a resident
        # can submit the required BOOKING_TOTAL payment while
        # the booking is still PENCIL.
        # ----------------------------------------------------

        financials = (
            FacilityBookingService.apply_financials(
                booking
            )
        )

        booking.rental_fee = financials[
            "rental_fee"
        ]

        booking.discount_amount = financials[
            "discount_amount"
        ]

        booking.security_deposit = financials[
            "security_deposit"
        ]

        booking.total_amount_due = financials[
            "total_amount_due"
        ]

        booking.save(
            update_fields=[
                "rental_fee",
                "discount_amount",
                "security_deposit",
                "total_amount_due",
                "updated_at",
            ]
        )

    # ========================================================
    # MINE
    # ========================================================

    @action(
        detail=False,
        methods=["get"],
        url_path="mine",
    )
    def mine(self, request):
        bookings = (
            self.get_queryset()
            .filter(
                resident__user=request.user
            )
            .order_by(
                "-booking_date",
                "-start_time",
            )
        )

        serializer = self.get_serializer(
            bookings,
            many=True,
        )

        return Response(
            serializer.data
        )

    # ========================================================
    # PENCIL / RENEW
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="pencil",
    )
    def pencil(self, request, pk=None):
        with transaction.atomic():

            booking = (
                self.get_queryset()
                .select_for_update()
                .select_related(
                    "facility",
                    "resident",
                    "resident__user",
                )
                .filter(pk=pk)
                .first()
            )

            if not booking:
                return Response(
                    {
                        "detail":
                            "Booking was not found."
                    },
                    status=status.HTTP_404_NOT_FOUND,
                )

            if (
                request.user.role
                != request.user.Roles.ADMIN
            ):
                if (
                    not booking.resident
                    or booking.resident.user_id
                    != request.user.id
                ):
                    raise PermissionDenied(
                        "You can only renew your own facility booking."
                    )

            if booking.status not in [
                Booking.Status.PENCIL,
                Booking.Status.PENDING,
            ]:
                return Response(
                    {
                        "detail":
                            (
                                "Only Pencil or Pending "
                                "bookings can be renewed."
                            )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            availability = (
                FacilityBookingService.check_availability(
                    facility=booking.facility,
                    booking_date=booking.booking_date,
                    start_time=booking.start_time,
                    end_time=booking.end_time,
                    estimated_guests=booking.estimated_guests,
                    exclude_booking_id=booking.id,
                )
            )

            if not availability.get(
                "available",
                False,
            ):
                return Response(
                    {
                        "detail":
                            availability.get(
                                "reason",
                                (
                                    "The selected facility "
                                    "time block is no longer available."
                                ),
                            )
                    },
                    status=status.HTTP_409_CONFLICT,
                )

            now = timezone.now()

            booking.status = Booking.Status.PENCIL
            booking.pencil_created_at = now
            booking.pencil_expires_at = (
                now
                + timedelta(
                    minutes=(
                        FacilityBookingService
                        .PENCIL_HOLD_MINUTES
                    )
                )
            )

            booking.save(
                update_fields=[
                    "status",
                    "pencil_created_at",
                    "pencil_expires_at",
                    "updated_at",
                ]
            )

        return Response(
            self.get_serializer(
                booking
            ).data
        )

    # ========================================================
    # SUBMIT
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="submit",
    )
    def submit(self, request, pk=None):

        with transaction.atomic():

            booking = (
                self.get_queryset()
                .select_for_update()
                .select_related(
                    "facility",
                    "resident",
                    "resident__user",
                )
                .prefetch_related(
                    "guests",
                    "payments",
                )
                .filter(pk=pk)
                .first()
            )

            if not booking:
                return Response(
                    {
                        "detail":
                            "Booking was not found."
                    },
                    status=status.HTTP_404_NOT_FOUND,
                )

            if booking.status != Booking.Status.PENCIL:
                return Response(
                    {
                        "detail":
                            (
                                "Only an active Pencil "
                                "booking can be submitted."
                            )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            valid, message = (
                FacilityBookingService
                .validate_active_pencil_booking(
                    booking
                )
            )

            if not valid:

                if booking.is_pencil_expired:

                    booking.status = (
                        Booking.Status.CANCELLED
                    )

                    booking.cancelled_at = (
                        timezone.now()
                    )

                    booking.cancellation_reason = (
                        "Pencil booking expired before submission."
                    )

                    booking.pencil_expires_at = None

                    booking.save(
                        update_fields=[
                            "status",
                            "cancelled_at",
                            "cancellation_reason",
                            "pencil_expires_at",
                            "updated_at",
                        ]
                    )

                    return Response(
                        {
                            "detail":
                                message
                        },
                        status=status.HTTP_409_CONFLICT,
                    )

                return Response(
                    {
                        "detail":
                            message
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            availability = (
                FacilityBookingService.check_availability(
                    facility=booking.facility,
                    booking_date=booking.booking_date,
                    start_time=booking.start_time,
                    end_time=booking.end_time,
                    estimated_guests=booking.estimated_guests,
                    exclude_booking_id=booking.id,
                )
            )

            if not availability.get(
                "available",
                False,
            ):
                return Response(
                    {
                        "detail":
                            availability.get(
                                "reason",
                                (
                                    "The selected facility "
                                    "time block is no longer available."
                                ),
                            )
                    },
                    status=status.HTTP_409_CONFLICT,
                )

            guest_count = (
                FacilityBookingService.guest_count(
                    booking
                )
            )

            guest_limit = int(
                booking.estimated_guests or 0
            )

            if guest_limit <= 0:
                return Response(
                    {
                        "detail":
                            (
                                "The booking does not contain "
                                "a valid guest limit."
                            )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            if guest_count <= 0:
                return Response(
                    {
                        "detail":
                            (
                                "At least one guest must be "
                                "added before submitting the booking."
                            )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            if guest_count > guest_limit:
                return Response(
                    {
                        "detail":
                            (
                                f"The guest list contains "
                                f"{guest_count} guests, exceeding "
                                f"the booking limit of "
                                f"{guest_limit}."
                            )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            # ------------------------------------------------
            # PAYMENT
            #
            # The amount is validated against the financial
            # snapshot already stored on the Pencil booking.
            # Do NOT re-price the booking here.
            # ------------------------------------------------

            payment = (
                booking.payments
                .filter(
                    payment_type=(
                        BookingPayment
                        .PaymentType
                        .BOOKING_TOTAL
                    ),
                    status__in=[
                        BookingPayment.Status.PENDING,
                        BookingPayment.Status.VERIFIED,
                    ],
                )
                .order_by(
                    "-created_at"
                )
                .first()
            )

            if payment:
                valid_payment, payment_message = (
                    FacilityBookingService
                    .validate_full_booking_payment(
                        booking,
                        payment.amount,
                    )
                )

                if not valid_payment:
                    return Response(
                        {
                            "detail":
                                payment_message
                        },
                        status=status.HTTP_400_BAD_REQUEST,
                    )

            elif (
                booking.total_amount_due
                > Decimal("0.00")
            ):
                return Response(
                    {
                        "detail":
                            (
                                "The full booking payment "
                                "must be submitted before "
                                "the reservation can be submitted."
                            )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            booking.status = Booking.Status.PENDING
            booking.submitted_at = timezone.now()
            booking.pencil_expires_at = None

            booking.save(
                update_fields=[
                    "status",
                    "submitted_at",
                    "pencil_expires_at",
                    "updated_at",
                ]
            )

            resident_name = (
                booking.resident.full_name
                if booking.resident
                else "Resident"
            )

            BookingViewSet._notify_admins(
                title="Facility Booking Request",
                message=(
                    f"{resident_name} submitted a "
                    f"facility booking request for "
                    f"{booking.facility.name} "
                    f"on {booking.booking_date}."
                ),
                actor=request.user,
                metadata={
                    "event":
                        "FACILITY_BOOKING_REQUESTED",
                    "booking_id":
                        booking.id,
                    "facility_id":
                        booking.facility_id,
                    "facility_name":
                        booking.facility.name,
                    "resident_id":
                        booking.resident_id,
                    "booking_date":
                        str(booking.booking_date),
                    "start_time":
                        str(booking.start_time),
                    "end_time":
                        str(booking.end_time),
                    "guest_count":
                        guest_count,
                    "total_amount_due":
                        str(booking.total_amount_due),
                },
                priority=Notification.Priority.INFO,
            )

        return Response(
            self.get_serializer(
                booking
            ).data
        )

    # ========================================================
    # APPROVE
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="approve",
        permission_classes=[
            IsAuthenticated,
            IsAdmin,
        ],
    )
    def approve(self, request, pk=None):

        with transaction.atomic():

            booking = (
                Booking.objects
                .select_for_update()
                .select_related(
                    "facility",
                    "resident",
                    "resident__user",
                )
                .prefetch_related(
                    "payments",
                    "guests",
                )
                .filter(pk=pk)
                .first()
            )

            if not booking:
                return Response(
                    {
                        "detail":
                            "Booking was not found."
                    },
                    status=status.HTTP_404_NOT_FOUND,
                )

            if booking.status != Booking.Status.PENDING:
                return Response(
                    {
                        "detail":
                            "Only pending bookings can be approved."
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            availability = (
                FacilityBookingService.check_availability(
                    facility=booking.facility,
                    booking_date=booking.booking_date,
                    start_time=booking.start_time,
                    end_time=booking.end_time,
                    estimated_guests=booking.estimated_guests,
                    exclude_booking_id=booking.id,
                )
            )

            if not availability.get(
                "available",
                False,
            ):
                return Response(
                    {
                        "detail":
                            availability.get(
                                "reason",
                                "Booking is no longer available.",
                            )
                    },
                    status=status.HTTP_409_CONFLICT,
                )

            total_due = (
                FacilityBookingService.money(
                    booking.total_amount_due
                )
            )

            if total_due > Decimal("0.00"):

                payment = (
                    booking.payments
                    .filter(
                        payment_type=(
                            BookingPayment
                            .PaymentType
                            .BOOKING_TOTAL
                        ),
                        status=(
                            BookingPayment
                            .Status
                            .VERIFIED
                        ),
                    )
                    .order_by(
                        "-verified_at",
                        "-created_at",
                    )
                    .first()
                )

                if not payment:
                    return Response(
                        {
                            "detail":
                                (
                                    "The full booking payment "
                                    "must be verified before "
                                    "the booking can be approved."
                                )
                        },
                        status=status.HTTP_400_BAD_REQUEST,
                    )

                valid_payment, payment_message = (
                    FacilityBookingService
                    .validate_full_booking_payment(
                        booking,
                        payment.amount,
                    )
                )

                if not valid_payment:
                    return Response(
                        {
                            "detail":
                                payment_message
                        },
                        status=status.HTTP_400_BAD_REQUEST,
                    )

            guest_count = (
                FacilityBookingService.guest_count(
                    booking
                )
            )

            if guest_count <= 0:
                return Response(
                    {
                        "detail":
                            (
                                "A guest list is required "
                                "before approval."
                            )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            if guest_count > booking.estimated_guests:
                return Response(
                    {
                        "detail":
                            (
                                "The guest list exceeds "
                                "the booking guest limit."
                            )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            # ------------------------------------------------
            # Preserve the financial snapshot captured when
            # the Pencil booking was created.
            # ------------------------------------------------

            booking.status = Booking.Status.APPROVED
            booking.approved_by = request.user
            booking.approved_at = timezone.now()

            booking.save(
                update_fields=[
                    "status",
                    "approved_by",
                    "approved_at",
                    "updated_at",
                ]
            )

            BookingViewSet._notify_resident(
                booking=booking,
                actor=request.user,
                title="Facility Booking Approved",
                message=(
                    f"Your booking for "
                    f"{booking.facility.name} "
                    f"on {booking.booking_date} "
                    f"has been approved."
                ),
                metadata={
                    "event":
                        "FACILITY_BOOKING_APPROVED",
                    "booking_id":
                        booking.id,
                    "facility_id":
                        booking.facility_id,
                    "facility_name":
                        booking.facility.name,
                },
                priority=Notification.Priority.SUCCESS,
            )

        return Response(
            self.get_serializer(
                booking
            ).data
        )

    # ========================================================
    # REJECT
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="reject",
        permission_classes=[
            IsAuthenticated,
            IsAdmin,
        ],
    )
    def reject(self, request, pk=None):

        booking = self.get_object()

        if booking.status not in [
            Booking.Status.PENCIL,
            Booking.Status.PENDING,
        ]:
            return Response(
                {
                    "detail":
                        (
                            "This booking cannot be "
                            "rejected in its current state."
                        )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        reason = str(
            request.data.get(
                "reason",
                "",
            )
        ).strip()

        if not reason:
            reason = (
                "The facility booking request was rejected."
            )

        booking.status = Booking.Status.REJECTED
        booking.rejection_reason = reason
        booking.pencil_expires_at = None

        booking.save(
            update_fields=[
                "status",
                "rejection_reason",
                "pencil_expires_at",
                "updated_at",
            ]
        )

        BookingViewSet._notify_resident(
            booking=booking,
            actor=request.user,
            title="Facility Booking Rejected",
            message=(
                f"Your booking for "
                f"{booking.facility.name} "
                f"on {booking.booking_date} "
                f"was rejected."
            ),
            metadata={
                "event":
                    "FACILITY_BOOKING_REJECTED",
                "booking_id":
                    booking.id,
                "facility_id":
                    booking.facility_id,
                "reason":
                    reason,
            },
            priority=Notification.Priority.WARNING,
        )

        return Response(
            self.get_serializer(
                booking
            ).data
        )

    # ========================================================
    # CANCEL
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="cancel",
    )
    def cancel(self, request, pk=None):

        booking = self.get_object()

        if booking.status not in [
            Booking.Status.PENCIL,
            Booking.Status.PENDING,
            Booking.Status.APPROVED,
        ]:
            return Response(
                {
                    "detail":
                        "This booking cannot be cancelled."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if (
            request.user.role
            != request.user.Roles.ADMIN
        ):
            if (
                not booking.resident
                or booking.resident.user_id
                != request.user.id
            ):
                raise PermissionDenied(
                    "You can only cancel your own booking."
                )

        reason = str(
            request.data.get(
                "reason",
                "",
            )
        ).strip()

        if not reason:
            reason = "Booking cancelled by user."

        booking.status = Booking.Status.CANCELLED
        booking.cancelled_at = timezone.now()
        booking.cancellation_reason = reason
        booking.pencil_expires_at = None

        booking.save(
            update_fields=[
                "status",
                "cancelled_at",
                "cancellation_reason",
                "pencil_expires_at",
                "updated_at",
            ]
        )

        if (
            request.user.role
            == request.user.Roles.ADMIN
        ):

            BookingViewSet._notify_resident(
                booking=booking,
                actor=request.user,
                title="Facility Booking Cancelled",
                message=(
                    f"Your booking for "
                    f"{booking.facility.name} "
                    f"on {booking.booking_date} "
                    f"has been cancelled."
                ),
                metadata={
                    "event":
                        "FACILITY_BOOKING_CANCELLED",
                    "booking_id":
                        booking.id,
                    "facility_id":
                        booking.facility_id,
                    "reason":
                        reason,
                },
                priority=Notification.Priority.WARNING,
            )

        else:

            BookingViewSet._notify_admins(
                title="Facility Booking Cancelled",
                message=(
                    f"{booking.resident.full_name} "
                    f"cancelled the booking for "
                    f"{booking.facility.name}."
                ),
                actor=request.user,
                metadata={
                    "event":
                        "FACILITY_BOOKING_CANCELLED",
                    "booking_id":
                        booking.id,
                    "facility_id":
                        booking.facility_id,
                    "resident_id":
                        booking.resident_id,
                    "reason":
                        reason,
                },
                priority=Notification.Priority.WARNING,
            )

        return Response(
            self.get_serializer(
                booking
            ).data
        )

    # ========================================================
    # START USE
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="start-use",
        permission_classes=[
            IsAuthenticated,
            IsAdmin,
        ],
    )
    def start_use(self, request, pk=None):

        booking = self.get_object()

        if booking.status != Booking.Status.APPROVED:
            return Response(
                {
                    "detail":
                        "Only approved bookings can be started."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if (
            booking.security_clearance_required
            and not booking.security_clearance_completed
        ):
            return Response(
                {
                    "detail":
                        (
                            "Security clearance must be completed "
                            "before the facility can be used."
                        )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        booking.status = Booking.Status.IN_USE

        booking.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        BookingViewSet._notify_resident(
            booking=booking,
            actor=request.user,
            title="Facility Booking Started",
            message=(
                f"Your booking for "
                f"{booking.facility.name} "
                f"is now marked as in use."
            ),
            metadata={
                "event":
                    "FACILITY_BOOKING_STARTED",
                "booking_id":
                    booking.id,
                "facility_id":
                    booking.facility_id,
            },
            priority=Notification.Priority.INFO,
        )

        return Response(
            self.get_serializer(
                booking
            ).data
        )

    # ========================================================
    # COMPLETE
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="complete",
        permission_classes=[
            IsAuthenticated,
            IsAdmin,
        ],
    )
    def complete(self, request, pk=None):

        booking = self.get_object()

        if booking.status != Booking.Status.IN_USE:
            return Response(
                {
                    "detail":
                        (
                            "Only bookings currently "
                            "in use can be completed."
                        )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if booking.inspection_required:
            booking.status = (
                Booking.Status.INSPECTION_PENDING
            )
        else:
            booking.status = Booking.Status.CLOSED

        booking.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        if (
            booking.status
            == Booking.Status.INSPECTION_PENDING
        ):

            BookingViewSet._notify_resident(
                booking=booking,
                actor=request.user,
                title="Facility Booking Awaiting Inspection",
                message=(
                    f"Your {booking.facility.name} "
                    f"booking has ended and is awaiting "
                    f"post-event inspection."
                ),
                metadata={
                    "event":
                        "FACILITY_INSPECTION_PENDING",
                    "booking_id":
                        booking.id,
                    "facility_id":
                        booking.facility_id,
                },
                priority=Notification.Priority.INFO,
            )

        else:

            BookingViewSet._notify_resident(
                booking=booking,
                actor=request.user,
                title="Facility Booking Closed",
                message=(
                    f"Your {booking.facility.name} "
                    f"booking has been completed and closed."
                ),
                metadata={
                    "event":
                        "FACILITY_BOOKING_CLOSED",
                    "booking_id":
                        booking.id,
                    "facility_id":
                        booking.facility_id,
                },
                priority=Notification.Priority.SUCCESS,
            )

        return Response(
            self.get_serializer(
                booking
            ).data
        )

    # ========================================================
    # SECURITY CLEARANCE
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="clearance",
        permission_classes=[
            IsAuthenticated,
            IsAdmin,
        ],
    )
    def clearance(self, request, pk=None):

        booking = self.get_object()

        if booking.status in [
            Booking.Status.CANCELLED,
            Booking.Status.REJECTED,
            Booking.Status.CLOSED,
        ]:
            return Response(
                {
                    "detail":
                        (
                            "Security clearance cannot "
                            "be completed for this booking."
                        )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if booking.security_clearance_completed:
            return Response(
                {
                    "detail":
                        "Security clearance is already completed."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if (
            booking.security_clearance_required
            is False
        ):
            return Response(
                {
                    "detail":
                        "Security clearance is not required for this booking."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        guest_count = (
            FacilityBookingService.guest_count(
                booking
            )
        )

        if guest_count <= 0:
            return Response(
                {
                    "detail":
                        (
                            "A guest list must be completed "
                            "before security clearance."
                        )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if guest_count > booking.estimated_guests:
            return Response(
                {
                    "detail":
                        (
                            "The guest list exceeds "
                            "the booking guest limit."
                        )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        now = timezone.now()

        booking.security_clearance_completed = True
        booking.security_cleared_by = request.user
        booking.security_cleared_at = now

        booking.save(
            update_fields=[
                "security_clearance_completed",
                "security_cleared_by",
                "security_cleared_at",
                "updated_at",
            ]
        )

        BookingViewSet._notify_resident(
            booking=booking,
            actor=request.user,
            title="Facility Security Clearance Completed",
            message=(
                f"Security clearance for your "
                f"{booking.facility.name} booking "
                f"has been completed. "
                f"The guest list is now locked."
            ),
            metadata={
                "event":
                    "FACILITY_SECURITY_CLEARANCE_COMPLETED",
                "booking_id":
                    booking.id,
                "facility_id":
                    booking.facility_id,
                "cleared_by":
                    request.user.id,
                "cleared_at":
                    now.isoformat(),
            },
            priority=Notification.Priority.SUCCESS,
        )

        return Response(
            self.get_serializer(
                booking
            ).data
        )

    # ========================================================
    # INSPECTION
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="inspect",
        permission_classes=[
            IsAuthenticated,
            IsAdmin,
        ],
    )
    def inspect_booking(
        self,
        request,
        pk=None,
    ):

        booking = self.get_object()

        if booking.status != (
            Booking.Status.INSPECTION_PENDING
        ):
            return Response(
                {
                    "detail":
                        (
                            "This booking is not "
                            "awaiting inspection."
                        )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        result = request.data.get(
            "result"
        )

        notes = str(
            request.data.get(
                "notes",
                "",
            )
        ).strip()

        raw_deduction = request.data.get(
            "deduction_amount",
            0,
        )

        valid_results = [
            value
            for value, label
            in BookingInspection.Result.choices
        ]

        if result not in valid_results:
            return Response(
                {
                    "result":
                        "Invalid inspection result."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            deduction_amount = Decimal(
                str(
                    raw_deduction
                    if raw_deduction not in [
                        None,
                        "",
                    ]
                    else "0"
                )
            )
        except (
            InvalidOperation,
            TypeError,
            ValueError,
        ):
            return Response(
                {
                    "deduction_amount":
                        (
                            "Deduction amount must "
                            "be a valid number."
                        )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if deduction_amount < Decimal("0.00"):
            return Response(
                {
                    "deduction_amount":
                        "Deduction amount cannot be negative."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            refund_data = (
                FacilityBookingService.calculate_refund(
                    booking,
                    deduction_amount,
                )
            )
        except ValueError as exc:
            return Response(
                {
                    "detail":
                        str(exc)
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        refund_amount = refund_data[
            "refund_amount"
        ]

        deduction_amount = refund_data[
            "deduction_amount"
        ]

        security_deposit = refund_data[
            "security_deposit"
        ]

        with transaction.atomic():

            booking = (
                Booking.objects
                .select_for_update()
                .select_related(
                    "facility",
                    "resident",
                    "resident__user",
                )
                .get(
                    pk=booking.pk
                )
            )

            inspection, _ = (
                BookingInspection.objects
                .select_for_update()
                .get_or_create(
                    booking=booking,
                    defaults={
                        "result":
                            BookingInspection
                            .Result
                            .PENDING,
                    },
                )
            )

            inspection.inspected_by = request.user
            inspection.inspected_at = timezone.now()
            inspection.result = result
            inspection.notes = notes
            inspection.deduction_amount = (
                deduction_amount
            )
            inspection.refund_amount = (
                refund_amount
            )

            inspection.save()

            if (
                security_deposit
                > Decimal("0.00")
            ):
                booking.status = (
                    Booking.Status.REFUND_PENDING
                )
            else:
                booking.status = (
                    Booking.Status.CLOSED
                )

            booking.save(
                update_fields=[
                    "status",
                    "updated_at",
                ]
            )

            if result == (
                BookingInspection.Result.PASSED
            ):
                title = "Facility Inspection Passed"

                message = (
                    f"Your {booking.facility.name} "
                    f"booking passed the post-event inspection."
                )

                priority = (
                    Notification.Priority.SUCCESS
                )

            elif result == (
                BookingInspection.Result.DAMAGE_FOUND
            ):
                title = "Facility Inspection Completed"

                message = (
                    f"The post-event inspection for your "
                    f"{booking.facility.name} booking found "
                    f"damage requiring review."
                )

                priority = (
                    Notification.Priority.WARNING
                )

            elif result == (
                BookingInspection.Result.CLEANUP_REQUIRED
            ):
                title = "Facility Inspection Completed"

                message = (
                    f"The post-event inspection for your "
                    f"{booking.facility.name} booking "
                    f"requires cleanup."
                )

                priority = (
                    Notification.Priority.WARNING
                )

            else:
                title = "Facility Inspection Completed"

                message = (
                    f"The post-event inspection for your "
                    f"{booking.facility.name} booking "
                    f"has been completed."
                )

                priority = (
                    Notification.Priority.INFO
                )

            BookingViewSet._notify_resident(
                booking=booking,
                actor=request.user,
                title=title,
                message=message,
                metadata={
                    "event":
                        "FACILITY_INSPECTION_COMPLETED",
                    "booking_id":
                        booking.id,
                    "facility_id":
                        booking.facility_id,
                    "inspection_result":
                        result,
                    "deduction_amount":
                        str(deduction_amount),
                    "refund_amount":
                        str(refund_amount),
                },
                priority=priority,
            )

        return Response(
            self.get_serializer(
                booking
            ).data
        )

    # ========================================================
    # REFUND
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="refund",
        permission_classes=[
            IsAuthenticated,
            IsAdmin,
        ],
    )
    def refund(self, request, pk=None):

        with transaction.atomic():

            booking = (
                Booking.objects
                .select_for_update()
                .select_related(
                    "facility",
                    "resident",
                    "resident__user",
                )
                .get(pk=pk)
            )

            if booking.status != (
                Booking.Status.REFUND_PENDING
            ):
                return Response(
                    {
                        "detail":
                            (
                                "This booking is not "
                                "awaiting a security "
                                "deposit refund."
                            )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            inspection = getattr(
                booking,
                "inspection",
                None,
            )

            if not inspection:
                return Response(
                    {
                        "detail":
                            (
                                "A post-event inspection "
                                "is required before refunding "
                                "the deposit."
                            )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            refund_amount = (
                FacilityBookingService.money(
                    inspection.refund_amount
                )
            )

            payment_method = request.data.get(
                "payment_method",
                BookingPayment
                .PaymentMethod
                .CASH,
            )

            valid_methods = [
                value
                for value, label
                in BookingPayment.PaymentMethod.choices
            ]

            if payment_method not in valid_methods:
                return Response(
                    {
                        "payment_method":
                            "Invalid payment method."
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            reference_number = str(
                request.data.get(
                    "reference_number",
                    "",
                )
            ).strip()

            existing_refund = (
                BookingPayment.objects
                .filter(
                    booking=booking,
                    payment_type=(
                        BookingPayment
                        .PaymentType
                        .SECURITY_DEPOSIT_REFUND
                    ),
                    status=(
                        BookingPayment
                        .Status
                        .REFUNDED
                    ),
                )
                .first()
            )

            if existing_refund:
                return Response(
                    {
                        "detail":
                            (
                                "The security deposit refund "
                                "has already been processed."
                            )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            if refund_amount > Decimal("0.00"):

                BookingPayment.objects.create(
                    booking=booking,
                    payment_type=(
                        BookingPayment
                        .PaymentType
                        .SECURITY_DEPOSIT_REFUND
                    ),
                    payment_method=payment_method,
                    direction=(
                        BookingPayment
                        .Direction
                        .OUTGOING
                    ),
                    amount=refund_amount,
                    reference_number=reference_number,
                    status=(
                        BookingPayment
                        .Status
                        .REFUNDED
                    ),
                    paid_at=timezone.now(),
                    verified_at=timezone.now(),
                    verified_by=request.user,
                )

            booking.status = Booking.Status.CLOSED

            booking.save(
                update_fields=[
                    "status",
                    "updated_at",
                ]
            )

            BookingViewSet._notify_resident(
                booking=booking,
                actor=request.user,
                title="Security Deposit Refunded",
                message=(
                    f"Your security deposit refund "
                    f"for {booking.facility.name} "
                    f"is ₱{refund_amount:,.2f}."
                ),
                metadata={
                    "event":
                        "FACILITY_DEPOSIT_REFUNDED",
                    "booking_id":
                        booking.id,
                    "facility_id":
                        booking.facility_id,
                    "refund_amount":
                        str(refund_amount),
                    "payment_method":
                        payment_method,
                    "reference_number":
                        reference_number,
                },
                priority=Notification.Priority.SUCCESS,
            )

        return Response(
            self.get_serializer(
                booking
            ).data
        )

    @action(
    detail=True,
    methods=["post"],
    url_path="guest-list/import",
    parser_classes=[MultiPartParser, FormParser],
)
    def import_guest_list(self, request, pk=None):
        """
        Import a booking guest list from an XLSX file.

        Expected columns:
            Guest Name
            Guest Vehicle Plate Number
            Guest Vehicle Model

        The importer intentionally ignores:
            - title rows above the table
            - blank rows
            - privacy/instruction/footer rows below the table

        Existing guests are replaced only after the entire XLSX has
        been successfully validated.
        """

        from django.core.exceptions import ValidationError as DjangoValidationError
        from django.db import IntegrityError

        # ---------------------------------------------------------
        # 1. Get booking
        # ---------------------------------------------------------
        booking = self.get_queryset().select_related(
            "facility",
            "resident",
        ).filter(pk=pk).first()

        if not booking:
            raise ValidationError("Booking not found.")

        # ---------------------------------------------------------
        # 2. Ownership check
        # ---------------------------------------------------------
        if (
            booking.resident
            and hasattr(booking.resident, "user")
            and booking.resident.user != request.user
            and not request.user.is_staff
        ):
            raise PermissionDenied(
                "You do not have permission to modify this booking."
            )

        # ---------------------------------------------------------
        # 3. Security lock
        # ---------------------------------------------------------
        if booking.security_clearance_completed:
            raise ValidationError(
                "The guest list can no longer be modified because "
                "security clearance has already been completed."
            )

        # ---------------------------------------------------------
        # 4. Only editable booking statuses may import guests
        # ---------------------------------------------------------
        editable_statuses = {
            Booking.Status.PENCIL,
            Booking.Status.PENDING,
        }

        if booking.status not in editable_statuses:
            raise ValidationError(
                "The guest list can only be imported while the booking "
                "is in PENCIL or PENDING status."
            )

        # ---------------------------------------------------------
        # 5. Expire old pencil bookings if necessary
        # ---------------------------------------------------------
        if booking.status == Booking.Status.PENCIL:
            FacilityBookingService.expire_pencil_bookings()

            booking.refresh_from_db(
                fields=[
                    "status",
                    "pencil_expires_at",
                    "security_clearance_completed",
                    "estimated_guests",
                ]
            )

            if booking.status != Booking.Status.PENCIL:
                raise ValidationError(
                    "This booking hold has expired and can no longer "
                    "accept a guest list."
                )

            if (
                booking.pencil_expires_at
                and timezone.now() >= booking.pencil_expires_at
            ):
                raise ValidationError(
                    "This booking hold has expired and can no longer "
                    "accept a guest list."
                )

        # ---------------------------------------------------------
        # 6. Validate uploaded file
        # ---------------------------------------------------------
        uploaded_file = request.FILES.get("file")

        if not uploaded_file:
            uploaded_file = request.FILES.get("guest_list")

        if not uploaded_file:
            uploaded_file = request.FILES.get("xlsx")

        if not uploaded_file:
            raise ValidationError(
                "Please upload an XLSX guest-list file."
            )

        filename = str(uploaded_file.name or "").lower()

        if not filename.endswith(".xlsx"):
            raise ValidationError(
                "Only XLSX files are supported for guest-list import."
            )

        # ---------------------------------------------------------
        # 7. Load workbook
        # ---------------------------------------------------------
        try:
            workbook = load_workbook(
                uploaded_file,
                read_only=True,
                data_only=True,
            )
        except InvalidFileException:
            raise ValidationError(
                "The uploaded file is not a valid XLSX workbook."
            )
        except Exception:
            raise ValidationError(
                "The uploaded XLSX file could not be read."
            )

        try:
            if not workbook.worksheets:
                raise ValidationError(
                    "The uploaded workbook does not contain a worksheet."
                )

            # -----------------------------------------------------
            # 8. Find worksheet containing the guest-list headers
            # -----------------------------------------------------
            required_headers = {
                "guest name": "full_name",
                "guest vehicle plate number": "vehicle_plate_number",
                "guest vehicle model": "vehicle_model",
            }

            header_row_number = None
            header_columns = None
            worksheet = None

            for current_sheet in workbook.worksheets:

                for row_number, row in enumerate(
                    current_sheet.iter_rows(values_only=True),
                    start=1,
                ):
                    normalized_cells = []

                    for cell in row:
                        if cell is None:
                            normalized_cells.append("")
                        else:
                            normalized_cells.append(
                                " ".join(str(cell).strip().lower().split())
                            )

                    discovered_columns = {}

                    for column_index, normalized_value in enumerate(
                        normalized_cells
                    ):
                        if normalized_value in required_headers:
                            discovered_columns[
                                required_headers[normalized_value]
                            ] = column_index

                    if set(discovered_columns.keys()) == set(
                        required_headers.values()
                    ):
                        worksheet = current_sheet
                        header_row_number = row_number
                        header_columns = discovered_columns
                        break

                if worksheet is not None:
                    break

            if worksheet is None:
                raise ValidationError(
                    "The XLSX file must contain these columns: "
                    "Guest Name, Guest Vehicle Plate Number, "
                    "Guest Vehicle Model."
                )

            # -----------------------------------------------------
            # 9. Parse guest rows
            # -----------------------------------------------------
            rows_for_import = []

            for row_number, row in enumerate(
                worksheet.iter_rows(values_only=True),
                start=1,
            ):

                # Ignore everything above the discovered header.
                if row_number <= header_row_number:
                    continue

                def get_cell(column_index):
                    if column_index is None:
                        return ""

                    if column_index >= len(row):
                        return ""

                    value = row[column_index]

                    if value is None:
                        return ""

                    return str(value).strip()

                raw_name = get_cell(
                    header_columns["full_name"]
                )

                raw_plate = get_cell(
                    header_columns["vehicle_plate_number"]
                )

                raw_model = get_cell(
                    header_columns["vehicle_model"]
                )

                # -------------------------------------------------
                # Completely empty row
                # -------------------------------------------------
                if not raw_name and not raw_plate and not raw_model:
                    continue

                # -------------------------------------------------
                # Repeated header row
                # -------------------------------------------------
                if (
                    raw_name.strip().lower() == "guest name"
                    and raw_plate.strip().lower()
                    == "guest vehicle plate number"
                    and raw_model.strip().lower()
                    == "guest vehicle model"
                ):
                    continue

                # -------------------------------------------------
                # Ignore footer/privacy/instruction text.
                #
                # A legitimate guest name cannot exceed 200 chars.
                # If a >200-character value appears in the name
                # column with NO vehicle information, it is treated
                # as worksheet footer/instruction content rather
                # than a guest record.
                # -------------------------------------------------
                if len(raw_name) > 200:

                    if not raw_plate and not raw_model:
                        break

                    raise ValidationError(
                        f"Guest Name in worksheet row {row_number} "
                        "must not exceed 200 characters."
                    )

                # -------------------------------------------------
                # Ignore obvious worksheet footer/instruction rows
                # where the first cell contains a long textual notice.
                # -------------------------------------------------
                lower_name = raw_name.lower()

                footer_markers = (
                    "privacy notice",
                    "privacy policy",
                    "personal data",
                    "data privacy",
                    "confidential",
                    "important notice",
                    "important:",
                    "note:",
                    "instruction:",
                    "instructions:",
                    "do not modify",
                    "do not edit",
                    "generated by",
                )

                if (
                    raw_name
                    and not raw_plate
                    and not raw_model
                    and any(
                        marker in lower_name
                        for marker in footer_markers
                    )
                ):
                    break

                # -------------------------------------------------
                # Actual guest row
                # -------------------------------------------------
                if not raw_name:
                    raise ValidationError(
                        f"Guest Name is required in worksheet row "
                        f"{row_number}."
                    )

                if len(raw_name) > 200:
                    raise ValidationError(
                        f"Guest Name in worksheet row {row_number} "
                        "must not exceed 200 characters."
                    )

                if len(raw_plate) > 30:
                    raise ValidationError(
                        f"Guest Vehicle Plate Number in worksheet row "
                        f"{row_number} must not exceed 30 characters."
                    )

                if len(raw_model) > 150:
                    raise ValidationError(
                        f"Guest Vehicle Model in worksheet row "
                        f"{row_number} must not exceed 150 characters."
                    )

                rows_for_import.append(
                    {
                        "full_name": raw_name,
                        "vehicle_plate_number": raw_plate.upper(),
                        "vehicle_model": raw_model,
                    }
                )

            # -----------------------------------------------------
            # 10. Validate guest count
            # -----------------------------------------------------
            guest_limit = int(
                booking.estimated_guests or 0
            )

            if guest_limit < 1:
                raise ValidationError(
                    "This booking does not have a valid guest limit."
                )

            if len(rows_for_import) > guest_limit:
                raise ValidationError(
                    f"The uploaded guest list contains "
                    f"{len(rows_for_import)} guests, but this booking "
                    f"allows only {guest_limit} guests."
                )

            if not rows_for_import:
                raise ValidationError(
                    "No guest records were found in the uploaded XLSX file."
                )

            # -----------------------------------------------------
            # 11. Normalize and validate all rows BEFORE deleting
            #     the existing guest list.
            # -----------------------------------------------------
            normalized_rows = []

            for index, guest_data in enumerate(
                rows_for_import,
                start=1,
            ):
                full_name = " ".join(
                    str(guest_data["full_name"]).strip().split()
                )

                vehicle_plate_number = " ".join(
                    str(
                        guest_data.get(
                            "vehicle_plate_number",
                            "",
                        )
                    ).strip().split()
                ).upper()

                vehicle_model = " ".join(
                    str(
                        guest_data.get(
                            "vehicle_model",
                            "",
                        )
                    ).strip().split()
                )

                if not full_name:
                    raise ValidationError(
                        f"Guest Name is required for guest row {index}."
                    )

                if len(full_name) > 200:
                    raise ValidationError(
                        f"Guest Name for guest row {index} "
                        "must not exceed 200 characters."
                    )

                if len(vehicle_plate_number) > 30:
                    raise ValidationError(
                        f"Guest Vehicle Plate Number for guest row "
                        f"{index} must not exceed 30 characters."
                    )

                if len(vehicle_model) > 150:
                    raise ValidationError(
                        f"Guest Vehicle Model for guest row {index} "
                        "must not exceed 150 characters."
                    )

                normalized_rows.append(
                    {
                        "full_name": full_name,
                        "vehicle_plate_number": vehicle_plate_number,
                        "vehicle_model": vehicle_model,
                    }
                )

            # -----------------------------------------------------
            # 12. Atomic database replacement
            # -----------------------------------------------------
            with transaction.atomic():

                locked_booking = (
                    Booking.objects
                    .select_for_update()
                    .select_related(
                        "facility",
                        "resident",
                    )
                    .get(pk=booking.pk)
                )

                # Re-check security lock.
                if locked_booking.security_clearance_completed:
                    raise ValidationError(
                        "The guest list can no longer be modified because "
                        "security clearance has already been completed."
                    )

                # Re-check status.
                if locked_booking.status not in editable_statuses:
                    raise ValidationError(
                        "The guest list can only be imported while the "
                        "booking is in PENCIL or PENDING status."
                    )

                # Re-check pencil expiry while holding the database lock.
                if locked_booking.status == Booking.Status.PENCIL:
                    if (
                        locked_booking.pencil_expires_at
                        and timezone.now()
                        >= locked_booking.pencil_expires_at
                    ):
                        raise ValidationError(
                            "This booking hold has expired and can no "
                            "longer accept a guest list."
                        )

                # Re-check guest limit.
                locked_guest_limit = int(
                    locked_booking.estimated_guests or 0
                )

                if len(normalized_rows) > locked_guest_limit:
                    raise ValidationError(
                        f"The uploaded guest list contains "
                        f"{len(normalized_rows)} guests, but this booking "
                        f"allows only {locked_guest_limit} guests."
                    )

                # -------------------------------------------------
                # Build new guest objects first.
                #
                # This allows Django model validation to happen before
                # deleting the existing guest list.
                # -------------------------------------------------
                new_guests = []

                for row_number, guest_data in enumerate(
                    normalized_rows,
                    start=1,
                ):
                    guest = BookingGuest(
                        booking=locked_booking,
                        row_number=row_number,
                        full_name=guest_data["full_name"],
                        vehicle_plate_number=guest_data[
                            "vehicle_plate_number"
                        ],
                        vehicle_model=guest_data[
                            "vehicle_model"
                        ],
                    )

                    try:
                        guest.full_clean()
                    except DjangoValidationError as exc:
                        message_dict = getattr(
                            exc,
                            "message_dict",
                            None,
                        )

                        if message_dict:
                            messages = []

                            for field_messages in message_dict.values():
                                messages.extend(
                                    str(message)
                                    for message in field_messages
                                )

                            raise ValidationError(
                                f"Invalid guest row {row_number}: "
                                + " ".join(messages)
                            )

                        raise ValidationError(
                            f"Invalid guest row {row_number}: "
                            f"{exc}"
                        )

                    new_guests.append(guest)

                # -------------------------------------------------
                # Replace existing guest list only now.
                # -------------------------------------------------
                BookingGuest.objects.filter(
                    booking=locked_booking
                ).delete()

                try:
                    BookingGuest.objects.bulk_create(
                        new_guests
                    )
                except IntegrityError:
                    raise ValidationError(
                        "The guest list could not be imported because "
                        "one or more guest records conflict with the "
                        "booking guest-list rules."
                    )

                created_guests = list(
                    BookingGuest.objects.filter(
                        booking=locked_booking
                    ).order_by(
                        "row_number",
                        "id",
                    )
                )

            # -----------------------------------------------------
            # 13. Return imported guests
            # -----------------------------------------------------
            serializer = BookingGuestSerializer(
                created_guests,
                many=True,
                context={
                    "request": request,
                },
            )

            return Response(
                {
                    "success": True,
                    "message": (
                        f"{len(created_guests)} guest record(s) "
                        "imported successfully."
                    ),
                    "guest_count": len(created_guests),
                    "guests": serializer.data,
                },
                status=status.HTTP_200_OK,
            )

        finally:
            try:
                workbook.close()
            except Exception:
                pass


# ============================================================
# BOOKING GUESTS
# ============================================================

class BookingGuestViewSet(ModelViewSet):
    """
    Guest-list records.

    Canonical endpoints:

        GET
        /api/facilities/booking-guests/

        POST
        /api/facilities/booking-guests/

        PATCH
        /api/facilities/booking-guests/<id>/

        DELETE
        /api/facilities/booking-guests/<id>/

    Resident ownership is enforced on the backend.

    Guest records become permanently locked after security
    clearance.
    """

    queryset = (
        BookingGuest.objects
        .select_related(
            "booking",
            "booking__facility",
            "booking__resident",
            "booking__resident__user",
        )
        .all()
    )

    serializer_class = BookingGuestSerializer

    permission_classes = [
        IsAuthenticated,
        IsAdminOrResident,
    ]

    http_method_names = [
        "get",
        "post",
        "patch",
        "delete",
        "head",
        "options",
    ]

    # ========================================================
    # QUERYSET VISIBILITY
    # ========================================================

    def get_queryset(self):
        queryset = super().get_queryset()

        user = self.request.user

        if user.role == user.Roles.ADMIN:
            return queryset

        return queryset.filter(
            booking__resident__user=user
        )

    # ========================================================
    # BOOKING ACCESS
    # ========================================================

    def _get_booking_for_request(
        self,
        booking_id,
    ):
        booking = (
            Booking.objects
            .select_related(
                "facility",
                "resident",
                "resident__user",
            )
            .filter(pk=booking_id)
            .first()
        )

        if not booking:
            raise ValidationError(
                {
                    "booking":
                        "Booking was not found."
                }
            )

        user = self.request.user

        if user.role != user.Roles.ADMIN:

            if (
                not booking.resident
                or booking.resident.user_id
                != user.id
            ):
                raise PermissionDenied(
                    "You can only manage guests for your own booking."
                )

        if booking.security_clearance_completed:
            raise ValidationError(
                {
                    "booking":
                        (
                            "The guest list is locked because "
                            "security clearance has been completed."
                        )
                }
            )

        if booking.status not in [
            Booking.Status.PENCIL,
            Booking.Status.PENDING,
            Booking.Status.APPROVED,
        ]:
            raise ValidationError(
                {
                    "booking":
                        (
                            "Guests can only be managed "
                            "while the booking is active."
                        )
                }
            )

        if (
            booking.status == Booking.Status.PENCIL
            and booking.is_pencil_expired
        ):
            raise ValidationError(
                {
                    "booking":
                        (
                            "This Pencil Book has expired. "
                            "Please create a new facility reservation."
                        )
                }
            )

        return booking

    # ========================================================
    # CREATE
    # ========================================================

    def perform_create(self, serializer):

        booking = serializer.validated_data[
            "booking"
        ]

        self._get_booking_for_request(
            booking.id
        )

        existing_count = (
            FacilityBookingService
            .guest_count(
                booking
            )
        )

        guest_limit = int(
            booking.estimated_guests or 0
        )

        if guest_limit <= 0:
            raise ValidationError(
                {
                    "booking":
                        (
                            "This booking does not have "
                            "a valid guest limit."
                        )
                }
            )

        if existing_count >= guest_limit:
            raise ValidationError(
                {
                    "booking":
                        (
                            f"The guest limit of "
                            f"{guest_limit} has already been reached."
                        )
                }
            )

        row_number = (
            FacilityBookingService
            .next_guest_row_number(
                booking
            )
        )

        serializer.save(
            row_number=row_number
        )

    # ========================================================
    # UPDATE
    # ========================================================

    def perform_update(self, serializer):

        guest = self.get_object()

        self._get_booking_for_request(
            guest.booking_id
        )

        serializer.save(
            row_number=guest.row_number
        )

    # ========================================================
    # DELETE
    # ========================================================

    def perform_destroy(self, instance):

        self._get_booking_for_request(
            instance.booking_id
        )

        instance.delete()

    # ========================================================
    # REORDER
    # ========================================================

    @action(
        detail=False,
        methods=["post"],
        url_path="reorder",
    )
    def reorder(self, request):

        booking_id = request.data.get(
            "booking"
        )

        guest_ids = request.data.get(
            "guest_ids"
        )

        if not booking_id:
            return Response(
                {
                    "booking":
                        "Booking is required."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not isinstance(
            guest_ids,
            list,
        ):
            return Response(
                {
                    "guest_ids":
                        "guest_ids must be a list."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        booking = self._get_booking_for_request(
            booking_id
        )

        guests = list(
            BookingGuest.objects
            .filter(
                booking=booking,
                id__in=guest_ids,
            )
        )

        if len(guests) != len(guest_ids):
            return Response(
                {
                    "guest_ids":
                        (
                            "One or more guest IDs do not "
                            "belong to this booking."
                        )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        guest_map = {
            guest.id: guest
            for guest in guests
        }

        with transaction.atomic():

            for row_number, guest_id in enumerate(
                guest_ids,
                start=1,
            ):

                guest = guest_map[
                    guest_id
                ]

                guest.row_number = row_number

                guest.save(
                    update_fields=[
                        "row_number",
                        "updated_at",
                    ]
                )

        return Response(
            {
                "detail":
                    "Guest list order updated successfully."
            }
        )


# ============================================================
# BOOKING PAYMENTS
# ============================================================

class BookingPaymentViewSet(ModelViewSet):
    """
    Facility booking payments.

    Resident:
        BOOKING_TOTAL

    Admin:
        verify / reject incoming booking payments
        and view all payment records.
    """

    queryset = (
        BookingPayment.objects
        .select_related(
            "booking",
            "booking__facility",
            "booking__resident",
            "booking__resident__user",
            "verified_by",
        )
        .all()
    )

    serializer_class = BookingPaymentSerializer

    permission_classes = [
        IsAuthenticated,
        IsAdminOrResident,
    ]

    parser_classes = [
        MultiPartParser,
        FormParser,
        JSONParser,
    ]

    http_method_names = [
        "get",
        "post",
        "patch",
        "head",
        "options",
    ]

    # ========================================================
    # QUERYSET VISIBILITY
    # ========================================================

    def get_queryset(self):

        queryset = super().get_queryset()

        user = self.request.user

        if user.role == user.Roles.ADMIN:
            return queryset

        return queryset.filter(
            booking__resident__user=user
        )

    # ========================================================
    # CREATE
    # ========================================================

    def perform_create(self, serializer):

        booking = serializer.validated_data[
            "booking"
        ]

        user = self.request.user

        if user.role != user.Roles.ADMIN:

            if (
                not booking.resident
                or booking.resident.user_id
                != user.id
            ):
                raise PermissionDenied(
                    "You can only submit payments for your own booking."
                )

        payment_type = serializer.validated_data.get(
            "payment_type"
        )

        amount = serializer.validated_data.get(
            "amount"
        )

        if (
            payment_type
            == BookingPayment
            .PaymentType
            .BOOKING_TOTAL
        ):

            if booking.status not in [
                Booking.Status.PENCIL,
                Booking.Status.PENDING,
            ]:
                raise ValidationError(
                    {
                        "booking":
                            (
                                "The booking cannot accept "
                                "a booking-total payment in "
                                "its current status."
                            )
                    }
                )

            valid, message = (
                FacilityBookingService
                .validate_full_booking_payment(
                    booking,
                    amount,
                )
            )

            if not valid:
                raise ValidationError(
                    {
                        "amount":
                            message
                    }
                )

        elif (
            payment_type
            == BookingPayment
            .PaymentType
            .SECURITY_DEPOSIT_REFUND
        ):

            if user.role != user.Roles.ADMIN:
                raise PermissionDenied(
                    "Only administrators can create security deposit refunds."
                )

        else:
            raise ValidationError(
                {
                    "payment_type":
                        "Invalid facility payment type."
                }
            )

        payment = serializer.save()

        if user.role != user.Roles.ADMIN:

            BookingViewSet._notify_admins(
                title="Facility Payment Submitted",
                message=(
                    f"{payment.booking.resident.full_name} "
                    f"submitted a ₱"
                    f"{payment.amount:,.2f} "
                    f"facility payment for "
                    f"{payment.booking.facility.name}."
                ),
                actor=user,
                metadata={
                    "event":
                        "FACILITY_PAYMENT_SUBMITTED",
                    "booking_id":
                        payment.booking_id,
                    "payment_id":
                        payment.id,
                    "amount":
                        str(payment.amount),
                    "payment_method":
                        payment.payment_method,
                    "payment_type":
                        payment.payment_type,
                },
                priority=Notification.Priority.INFO,
            )

    # ========================================================
    # VERIFY
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="verify",
        permission_classes=[
            IsAuthenticated,
            IsAdmin,
        ],
    )
    def verify(self, request, pk=None):

        with transaction.atomic():

            payment = (
                BookingPayment.objects
                .select_for_update()
                .select_related(
                    "booking",
                    "booking__facility",
                    "booking__resident",
                    "booking__resident__user",
                )
                .filter(pk=pk)
                .first()
            )

            if not payment:
                return Response(
                    {
                        "detail":
                            "Payment was not found."
                    },
                    status=status.HTTP_404_NOT_FOUND,
                )

            if payment.status != (
                BookingPayment.Status.PENDING
            ):
                return Response(
                    {
                        "detail":
                            "Only pending payments can be verified."
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            if payment.payment_type != (
                BookingPayment
                .PaymentType
                .BOOKING_TOTAL
            ):
                return Response(
                    {
                        "detail":
                            (
                                "Only booking-total "
                                "payments can be verified."
                            )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            booking = payment.booking

            valid, message = (
                FacilityBookingService
                .validate_full_booking_payment(
                    booking,
                    payment.amount,
                )
            )

            if not valid:
                return Response(
                    {
                        "detail":
                            message
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            payment.status = (
                BookingPayment.Status.VERIFIED
            )

            payment.verified_at = timezone.now()
            payment.verified_by = request.user

            payment.save(
                update_fields=[
                    "status",
                    "verified_at",
                    "verified_by",
                    "updated_at",
                ]
            )

            BookingViewSet._notify_resident(
                booking=booking,
                actor=request.user,
                title="Facility Payment Verified",
                message=(
                    f"Your ₱{payment.amount:,.2f} "
                    f"payment for "
                    f"{booking.facility.name} "
                    f"has been verified."
                ),
                metadata={
                    "event":
                        "FACILITY_PAYMENT_VERIFIED",
                    "booking_id":
                        booking.id,
                    "payment_id":
                        payment.id,
                    "amount":
                        str(payment.amount),
                    "payment_type":
                        payment.payment_type,
                },
                priority=Notification.Priority.SUCCESS,
            )

        return Response(
            self.get_serializer(
                payment
            ).data
        )

    # ========================================================
    # REJECT
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="reject",
        permission_classes=[
            IsAuthenticated,
            IsAdmin,
        ],
    )
    def reject(self, request, pk=None):

        payment = self.get_object()

        if payment.status != (
            BookingPayment.Status.PENDING
        ):
            return Response(
                {
                    "detail":
                        "Only pending payments can be rejected."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if payment.payment_type != (
            BookingPayment
            .PaymentType
            .BOOKING_TOTAL
        ):
            return Response(
                {
                    "detail":
                        (
                            "Only booking-total "
                            "payments can be rejected."
                        )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        payment.status = (
            BookingPayment.Status.REJECTED
        )

        payment.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        BookingViewSet._notify_resident(
            booking=payment.booking,
            actor=request.user,
            title="Facility Payment Rejected",
            message=(
                f"Your ₱{payment.amount:,.2f} "
                f"payment for "
                f"{payment.booking.facility.name} "
                f"was rejected. "
                f"Please submit a valid payment proof."
            ),
            metadata={
                "event":
                    "FACILITY_PAYMENT_REJECTED",
                "booking_id":
                    payment.booking_id,
                "payment_id":
                    payment.id,
                "amount":
                    str(payment.amount),
            },
            priority=Notification.Priority.WARNING,
        )

        return Response(
            self.get_serializer(
                payment
            ).data
        )


# ============================================================
# BOOKING INSPECTIONS
# ============================================================

class BookingInspectionViewSet(ModelViewSet):
    """
    Read-only inspection records.

    Inspection creation and booking state transitions are
    performed through BookingViewSet.inspect_booking().
    """

    queryset = (
        BookingInspection.objects
        .select_related(
            "booking",
            "booking__facility",
            "booking__resident",
            "booking__resident__user",
            "inspected_by",
        )
        .all()
    )

    serializer_class = BookingInspectionSerializer

    permission_classes = [
        IsAuthenticated,
        IsAdmin,
    ]

    http_method_names = [
        "get",
        "head",
        "options",
    ]

    def get_queryset(self):

        queryset = super().get_queryset()

        booking_id = (
            self.request.query_params.get(
                "booking"
            )
        )

        if booking_id:
            queryset = queryset.filter(
                booking_id=booking_id
            )

        return queryset
